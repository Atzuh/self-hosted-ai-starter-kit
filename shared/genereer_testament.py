#!/usr/bin/env python3
"""Bouwt concept-testamenten en de bijbehorende offerte uit tekstblokken.

    python3 /data/shared/genereer_testament.py --args-file /tmp/n8n_args_testament_XYZ.json

De payload bevat alleen de KEUZES, niet de opgeloste teksten: welke blokken mee
moeten, wat er in de invulvelden staat en wie de erflaters zijn. De blokteksten,
de tarieven en de kantoorgegevens leest dit script zelf van schijf. Zo staat de
prijsberekening op één plek — in Python, waar hij te testen is — in plaats van in
een Code node in een JSON-bestand.

Bij een echtpaar worden twee testamenten gemaakt. Die zijn spiegelbeeldig maar
niet noodzakelijk gelijk: elk testament heeft een eigen blokselectie en eigen
veldwaarden, want wat de een regelt hoeft de ander niet te regelen.

Wat er uit komt is een DOSSIERSTUK, geen akte. Dat staat in de kop van elke
pagina en in de voettekst, en alles wat niet is ingevuld staat geel gearceerd.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor

# document_to_preview_html werkt op elk python-docx Document en voedt de
# inline preview in de webapp. Het script staat in dezelfde map, en Python zet
# de scriptmap zelf op sys.path.
from vul_template_in import document_to_preview_html

OUTPUT_DIR = Path("/data/shared/output")
BLOKKEN_BESTAND = Path("/data/shared/templates/aktes/testament.json")
KANTOOR_BESTAND = Path("/data/shared/kantoor.json")

DEFAULT_FONT = "Calibri"
HEADING_COLOR = RGBColor(0x00, 0x00, 0x00)
INK_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
META_COLOR = RGBColor(0x73, 0x73, 0x73)

PLACEHOLDER_RE = re.compile(r"<<[A-Z0-9_]+>>")


# ---------------------------------------------------------------------------
# Opmaak-hulpjes (zelfde familie als genereer_besprekingsverslag.py)
# ---------------------------------------------------------------------------

def _set_run(run, *, bold=False, size=11, color=INK_COLOR, italic=False,
             gearceerd=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = DEFAULT_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if gearceerd:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_paragraph(doc: Document, text: str, *, bold=False, size=11,
                   color=INK_COLOR, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        _set_run(p.add_run(text), bold=bold, size=size, color=color, italic=italic)
    return p


def _add_heading(doc: Document, text: str, *, level: int = 1) -> None:
    sizes = {1: 16, 2: 12, 3: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    p.paragraph_format.space_after = Pt(4)
    _set_run(p.add_run(text), bold=True, size=sizes.get(level, 11), color=HEADING_COLOR)


def _sanitize(naam: str) -> str:
    schoon = re.sub(r"[^0-9A-Za-z_-]", "_", naam).strip("_")
    return schoon or "onbekend"


def _achternaam(volledige_naam: str) -> str:
    """Laatste woord van de naam, voor in de bestandsnaam.

    Een mappenlijst met 'concept_testament_2026-114_1.docx' en '_2.docx' zegt
    niet wiens testament het is; met de achternaam erin wel.
    """
    delen = [d for d in re.split(r"\s+", volledige_naam.strip()) if d]
    return _sanitize(delen[-1]) if delen else ""


def _euro(cent: int) -> str:
    heel, rest = divmod(abs(int(cent)), 100)
    teken = "-" if cent < 0 else ""
    duizend = f"{heel:,}".replace(",", ".")
    return f"{teken}€ {duizend},{rest:02d}"


# ---------------------------------------------------------------------------
# Tekstblokken invullen
# ---------------------------------------------------------------------------

def _persoon_zin(persoon: dict) -> str:
    """Comparitieregel voor één erflater, uit de ingevoerde gegevens.

    Ontbrekende onderdelen laten we weg in plaats van er lege komma's van te
    maken; wat helemaal ontbreekt wordt hierboven al als leeg veld gearceerd.
    """
    aanhef = str(persoon.get("aanhef") or "").strip()
    naam = str(persoon.get("naam_voluit") or "").strip()
    delen = [f"{aanhef} {naam}".strip() or ""]
    plaats = str(persoon.get("geboorteplaats") or "").strip()
    datum = str(persoon.get("geboortedatum") or "").strip()
    if plaats and datum:
        delen.append(f"geboren te {plaats} op {datum}")
    elif plaats:
        delen.append(f"geboren te {plaats}")
    elif datum:
        delen.append(f"geboren op {datum}")
    adres = str(persoon.get("adres") or "").strip()
    if adres:
        delen.append(f"wonende te {adres}")
    return ", ".join(d for d in delen if d)


def _schrijf_regel(doc: Document, regel: str, waarden: dict, labels: dict,
                   leeg_verzamelaar: set) -> None:
    """Zet één regel bloktekst om in een alinea, met gele gaten waar nodig.

    Een placeholder zonder waarde wordt «label» in geel. Dat is leesbaarder in
    Word dan de kale <<VOOGD_NAAM>>, en nog steeds te vinden door op « te zoeken.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    positie = 0
    for m in PLACEHOLDER_RE.finditer(regel):
        if m.start() > positie:
            _set_run(p.add_run(regel[positie:m.start()]))
        sleutel = m.group(0)[2:-2]
        waarde = str(waarden.get(sleutel, "") or "").strip()
        if waarde:
            _set_run(p.add_run(waarde))
        else:
            leeg_verzamelaar.add(sleutel)
            label = labels.get(sleutel) or sleutel.replace("_", " ").lower()
            _set_run(p.add_run(f"«{label}»"), gearceerd=True)
        positie = m.end()
    if positie < len(regel):
        _set_run(p.add_run(regel[positie:]))
    if not p.runs:
        _set_run(p.add_run(""))


# ---------------------------------------------------------------------------
# Het testament
# ---------------------------------------------------------------------------

def render_testament(*, blokken: list, kantoor: dict, erflater: dict,
                     partner: dict | None, velden: dict, referentie: str,
                     metadata: dict, herkomst: dict) -> tuple[Document, set]:
    doc = Document()
    sectie = doc.sections[0]
    sectie.left_margin = Cm(2.5)
    sectie.right_margin = Cm(2.5)
    sectie.top_margin = Cm(2.2)
    sectie.bottom_margin = Cm(2.2)

    # Kop- en voettekst op elke pagina. Een uit blokken samengesteld stuk mag
    # nergens voor een akte kunnen doorgaan, ook niet als iemand pagina 4
    # los uitprint.
    kop = sectie.header.paragraphs[0]
    kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(kop.add_run("CONCEPT — NIET VOOR ONDERTEKENING"), bold=True, size=9,
             color=META_COLOR)
    voet = sectie.footer.paragraphs[0]
    voet.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dossier = str(metadata.get("dossier") or referentie)
    _set_run(
        voet.add_run(f"Concept · dossier {dossier} · automatisch samengesteld uit tekstblokken"),
        size=8, color=META_COLOR, italic=True,
    )

    _add_paragraph(doc, "CONCEPT-TESTAMENT", bold=True, size=9, color=META_COLOR,
                   space_after=2)
    naam = str(erflater.get("naam_voluit") or "").strip()
    _add_heading(doc, f"Uiterste wilsbeschikking{' van ' + naam if naam else ''}", level=1)

    gemaakt_op = str(metadata.get("datum") or "").strip()
    toelichting = (
        "Dit stuk is samengesteld uit standaard tekstblokken op basis van het "
        "besprekingsverslag"
        + (f" van {herkomst['verslag_datum']}" if herkomst.get("verslag_datum") else "")
        + ". De samenhang tussen de blokken is niet gecontroleerd en de notaris stelt "
        "de definitieve tekst vast. Geel gearceerde plekken moeten nog worden ingevuld."
    )
    _add_paragraph(doc, toelichting, italic=True, size=9, color=META_COLOR, space_after=4)
    regel = f"Dossier {dossier}" + (f" · {gemaakt_op}" if gemaakt_op else "")
    _add_paragraph(doc, regel, size=9, color=META_COLOR, space_after=12)

    # Systeemvelden: de erflater van dít testament, en de kantoorgegevens. Dit
    # deel spiegelt zichzelf — in haar testament is zij de comparant.
    systeem = {
        "NOTARIS_NAAM": str(kantoor.get("notaris_naam") or ""),
        "NOTARIS_STANDPLAATS": str(kantoor.get("notaris_standplaats") or ""),
        "ERFLATERS": _persoon_zin(erflater),
        "PARTNER": _persoon_zin(partner) if partner else "",
        # Bewust leeg: de passeerdatum staat pas bij het passeren vast, en komt
        # dus gearceerd in het concept — net als in de hypotheekakte-flow.
        "AKTE_DATUM": "",
    }
    alle_waarden = {**systeem, **{k: v for k, v in velden.items()}}

    labels = {}
    for blok in blokken:
        for veld in blok.get("velden") or []:
            labels[str(veld.get("naam"))] = str(veld.get("label") or "")
    labels.setdefault("AKTE_DATUM", "datum van passeren")
    labels.setdefault("NOTARIS_NAAM", "naam notaris")
    labels.setdefault("NOTARIS_STANDPLAATS", "standplaats notaris")
    labels.setdefault("ERFLATERS", "gegevens van de erflater")

    slot_regels = kantoor.get("slot_tekst") or []
    leeg: set = set()

    for blok in blokken:
        kop_in_akte = str(blok.get("naam") or "")
        tekstregels = str(blok.get("tekst") or "").split("\n")
        eerste = True
        for regel_tekst in tekstregels:
            if regel_tekst.strip() == "<<SLOT_TEKST>>":
                if slot_regels:
                    _add_heading(doc, "SLOT", level=2)
                    for s in slot_regels:
                        _schrijf_regel(doc, str(s), alle_waarden, labels, leeg)
                else:
                    leeg.add("SLOT_TEKST")
                    _schrijf_regel(doc, "<<SLOT_TEKST>>", alle_waarden, labels, leeg)
                eerste = False
                continue
            if eerste and kop_in_akte and regel_tekst.strip().isupper() and regel_tekst.strip():
                # De bloktekst begint zelf met een kopregel in hoofdletters;
                # die zetten we als kopje in plaats van als gewone alinea.
                _add_heading(doc, regel_tekst.strip(), level=2)
                eerste = False
                continue
            _schrijf_regel(doc, regel_tekst, alle_waarden, labels, leeg)
            eerste = False

    _add_paragraph(doc, "", space_after=8)
    _add_paragraph(
        doc,
        "Dit concept is automatisch samengesteld en dient ter voorbereiding van de "
        "bespreking en de akte. De behandelend notaris blijft verantwoordelijk voor "
        "inhoud, samenhang en vorm.",
        italic=True, size=8.5, color=META_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    return doc, leeg


# ---------------------------------------------------------------------------
# De offerte
# ---------------------------------------------------------------------------

def bereken_offerte(*, bestand: dict, kantoor: dict, testamenten: list) -> dict:
    """Honorarium, btw en verschotten, in hele centen.

    Blokken worden ÉÉN keer geteld, ook als ze in beide testamenten staan: het
    denkwerk is één keer gedaan. Het tweede testament heeft daarvoor een eigen
    vast tarief.
    """
    basistarief = int(bestand.get("basistarief_cent") or 0)
    prijzen = {str(b.get("id")): int(b.get("prijs_cent") or 0) for b in bestand.get("blokken") or []}
    namen = {str(b.get("id")): str(b.get("naam") or b.get("id")) for b in bestand.get("blokken") or []}

    gekozen: list[str] = []
    for t in testamenten:
        for blok_id in t.get("geselecteerd") or []:
            if blok_id not in gekozen:
                gekozen.append(str(blok_id))

    regels = [{"omschrijving": f"Basistarief {bestand.get('naam') or 'akte'}",
               "bedrag_cent": basistarief}]
    for blok_id in gekozen:
        bedrag = prijzen.get(blok_id, 0)
        if bedrag:
            regels.append({"omschrijving": namen.get(blok_id, blok_id), "bedrag_cent": bedrag})

    extra = max(0, len(testamenten) - 1)
    tarief_tweede = int(bestand.get("tarief_tweede_testament_cent") or 0)
    for i in range(extra):
        regels.append({
            "omschrijving": "Tweede testament partner" if extra == 1 else f"Extra testament {i + 2}",
            "bedrag_cent": tarief_tweede,
        })

    honorarium = sum(r["bedrag_cent"] for r in regels)

    # Integer-rekenen en zelf afronden: round() rondt in Python half naar even,
    # en floats geven hier zichtbare centen-verschillen.
    promille = int(round(float(kantoor.get("btw_percentage") or 0) * 10))
    btw = (honorarium * promille + 500) // 1000

    verschotten = []
    for v in bestand.get("verschotten") or []:
        bedrag = int(v.get("bedrag_cent") or 0)
        aantal = len(testamenten) if v.get("per_testament") else 1
        if bedrag and aantal:
            verschotten.append({
                "omschrijving": str(v.get("naam") or ""),
                "aantal": aantal,
                "bedrag_cent": bedrag * aantal,
            })
    verschotten_totaal = sum(v["bedrag_cent"] for v in verschotten)

    return {
        "regels": regels,
        "honorarium_cent": honorarium,
        "btw_percentage": float(kantoor.get("btw_percentage") or 0),
        "btw_cent": btw,
        "verschotten": verschotten,
        "verschotten_cent": verschotten_totaal,
        "totaal_cent": honorarium + btw + verschotten_totaal,
    }


def render_offerte(*, offerte: dict, kantoor: dict, personen: list,
                   referentie: str, metadata: dict) -> Document:
    doc = Document()
    sectie = doc.sections[0]
    sectie.left_margin = Cm(2.5)
    sectie.right_margin = Cm(2.5)
    sectie.top_margin = Cm(2.2)
    sectie.bottom_margin = Cm(2.2)

    _add_paragraph(doc, "OFFERTE", bold=True, size=9, color=META_COLOR, space_after=2)
    _add_heading(doc, "Kosten voor het opstellen van uw testament", level=1)

    dossier = str(metadata.get("dossier") or referentie)
    datum = str(metadata.get("datum") or "").strip()
    kop = f"Dossier {dossier}" + (f" · {datum}" if datum else "")
    _add_paragraph(doc, kop, size=9, color=META_COLOR, space_after=10)

    namen = [str(p.get("naam_voluit") or "").strip() for p in personen]
    namen = [n for n in namen if n]
    if namen:
        _add_paragraph(doc, "Voor: " + " en ".join(namen), space_after=12)

    _add_heading(doc, "Honorarium", level=2)
    for regel in offerte["regels"]:
        _regel_met_bedrag(doc, regel["omschrijving"], regel["bedrag_cent"])
    _regel_met_bedrag(doc, "Subtotaal honorarium", offerte["honorarium_cent"], vet=True)
    btw_label = f"Btw {offerte['btw_percentage']:g}%".replace(".", ",")
    _regel_met_bedrag(doc, btw_label, offerte["btw_cent"])

    if offerte["verschotten"]:
        _add_heading(doc, "Verschotten", level=2)
        _add_paragraph(
            doc,
            "Kosten die wij voor u voorschieten en zonder btw doorbelasten.",
            italic=True, size=9, color=META_COLOR, space_after=4,
        )
        for v in offerte["verschotten"]:
            label = v["omschrijving"] + (f" ({v['aantal']}×)" if v["aantal"] > 1 else "")
            _regel_met_bedrag(doc, label, v["bedrag_cent"])

    _add_paragraph(doc, "", space_after=4)
    _regel_met_bedrag(doc, "Totaal", offerte["totaal_cent"], vet=True, size=12)

    slot = kantoor.get("offerte_slottekst") or []
    if slot:
        dagen = int(kantoor.get("offerte_geldigheidsdagen") or 30)
        geldig_tot = (date.today() + timedelta(days=dagen)).strftime("%d-%m-%Y")
        _add_paragraph(doc, "", space_after=8)
        for regel in slot:
            _add_paragraph(doc, str(regel).replace("<<GELDIG_TOT>>", geldig_tot),
                           size=9, color=META_COLOR, space_after=4)
    return doc


def _regel_met_bedrag(doc: Document, omschrijving: str, cent: int, *, vet=False,
                      size=11) -> None:
    """Omschrijving links, bedrag rechts, met een rechtse tabstop.

    Bewust geen tabel: dan blijft document_to_preview_html bruikbaar, die
    alleen paragrafen rendert.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_ALIGN_PARAGRAPH.RIGHT)
    _set_run(p.add_run(omschrijving), bold=vet, size=size)
    _set_run(p.add_run("\t" + _euro(cent)), bold=vet, size=size)


# ---------------------------------------------------------------------------
# Aanroep
# ---------------------------------------------------------------------------

def _load_args_file(pad: str) -> dict:
    """Lees de payload en verwijder het bestand direct na inlezen.

    Zelfde patroon als de andere generatoren: n8n schrijft de argumenten naar
    een tijdelijk bestand zodat cliëntgegevens nooit door de shell gaan.
    """
    p = Path(pad)
    inhoud = p.read_text(encoding="utf-8")
    try:
        p.unlink()
    except OSError:
        pass
    payload = json.loads(inhoud)
    if not isinstance(payload, dict):
        raise SystemExit("ERROR: --args-file payload moet een JSON-object zijn.")
    return payload


def _lees_json(pad: Path) -> dict:
    return json.loads(pad.read_text(encoding="utf-8"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--args-file", required=True, dest="args_file")
    args = parser.parse_args()

    payload = _load_args_file(args.args_file)

    bestand = _lees_json(Path(payload.get("blokken_bestand") or BLOKKEN_BESTAND))
    try:
        kantoor = _lees_json(Path(payload.get("kantoor_bestand") or KANTOOR_BESTAND))
    except (OSError, json.JSONDecodeError):
        kantoor = {}

    referentie = str(payload.get("referentie") or "onbekend")
    metadata = payload.get("metadata") or {}
    herkomst = payload.get("herkomst") or {}
    personen = payload.get("personen") or []
    testamenten = payload.get("testamenten") or []
    if not testamenten:
        raise SystemExit("ERROR: payload bevat geen 'testamenten'.")
    if not personen:
        raise SystemExit("ERROR: payload bevat geen 'personen'.")

    blokken_op_id = {str(b.get("id")): b for b in bestand.get("blokken") or []}
    volgorde = [str(b.get("id")) for b in bestand.get("blokken") or []]

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    documenten = []
    alle_lege_velden: set = set()
    gebruikte_namen: set = set()

    for index, t in enumerate(testamenten):
        erflater_index = int(t.get("erflater", index))
        if erflater_index >= len(personen):
            raise SystemExit(f"ERROR: testament {index + 1} verwijst naar persoon {erflater_index}.")
        erflater = personen[erflater_index]
        partner = next(
            (p for i, p in enumerate(personen) if i != erflater_index), None
        )

        gekozen = [bid for bid in volgorde if bid in set(t.get("geselecteerd") or [])]
        onbekend = sorted(set(t.get("geselecteerd") or []) - set(volgorde))
        if onbekend:
            raise SystemExit(f"ERROR: onbekende blokken: {', '.join(onbekend)}")
        blokken = [blokken_op_id[bid] for bid in gekozen]

        doc, leeg = render_testament(
            blokken=blokken, kantoor=kantoor, erflater=erflater, partner=partner,
            velden=t.get("velden") or {}, referentie=referentie, metadata=metadata,
            herkomst=herkomst,
        )
        alle_lege_velden |= leeg

        achter = _achternaam(str(erflater.get("naam_voluit") or ""))
        stam = f"concept_testament_{_sanitize(referentie)}"
        naam = f"{stam}_{achter}" if achter else f"{stam}_{index + 1}"
        if naam in gebruikte_namen:
            # Twee erflaters met dezelfde achternaam is eerder regel dan
            # uitzondering bij een echtpaar; dan het volgnummer erachter.
            naam = f"{naam}_{index + 1}"
        gebruikte_namen.add(naam)

        pad = OUTPUT_DIR / f"{naam}.docx"
        doc.save(str(pad))
        (OUTPUT_DIR / f"{naam}.html").write_text(document_to_preview_html(doc), encoding="utf-8")
        documenten.append({
            "soort": "akte",
            "titel": f"Concept-testament {erflater.get('naam_voluit') or index + 1}".strip(),
            "pad": str(pad),
            "bestandsnaam": pad.name,
            "blokken": len(blokken),
            "lege_velden": sorted(leeg),
        })

    offerte = bereken_offerte(bestand=bestand, kantoor=kantoor, testamenten=testamenten)
    doc = render_offerte(offerte=offerte, kantoor=kantoor, personen=personen,
                         referentie=referentie, metadata=metadata)
    offerte_naam = f"offerte_{_sanitize(referentie)}"
    offerte_pad = OUTPUT_DIR / f"{offerte_naam}.docx"
    doc.save(str(offerte_pad))
    (OUTPUT_DIR / f"{offerte_naam}.html").write_text(document_to_preview_html(doc), encoding="utf-8")
    documenten.append({
        "soort": "offerte",
        "titel": "Offerte",
        "pad": str(offerte_pad),
        "bestandsnaam": offerte_pad.name,
    })

    # De browser rekent hetzelfde uit; wijkt het af, dan is het blokkenbestand
    # tussen laden en genereren gewijzigd. Het document toont altijd de nieuwe
    # prijs; de melding gaat mee terug zodat de webapp het kan tonen.
    controle = payload.get("prijs_controle_cent")
    waarschuwingen = []
    if controle is not None and int(controle) != offerte["totaal_cent"]:
        waarschuwingen.append(
            f"Het totaalbedrag is {_euro(offerte['totaal_cent'])} en niet "
            f"{_euro(int(controle))} zoals in het scherm stond. "
            "Waarschijnlijk zijn de tarieven intussen gewijzigd."
        )

    print("DOCUMENTS: " + json.dumps(documenten, ensure_ascii=False))
    print("OFFERTE: " + json.dumps(offerte, ensure_ascii=False))
    print(f"TOTAAL_CENT: {offerte['totaal_cent']}")
    print(f"LEGE_VELDEN: {len(alle_lege_velden)}")
    for w in waarschuwingen:
        print("WAARSCHUWING: " + w)
    # De SUCCESS-regel wijst naar het eerste concept; alle Build-Response-nodes
    # in dit project zoeken op dat patroon.
    print(f"SUCCESS: {documenten[0]['pad']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
